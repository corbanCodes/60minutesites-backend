"""Shared plumbing: the customer's OpenAI key, spreadsheets in and out, and
what a run will cost before it is started."""
import csv
import io
import json
import math
import os
import re

import requests as http

from enrich.models import DEFAULT_MODEL, MODEL_BY_ID, MODELS

OPENAI = "https://api.openai.com/v1"


# ------------------------------------------------------------------ openai
def verify_key(key):
    """-> {ok, models:[ids], error}. Called when someone pastes a key, so the
    model picker only ever offers models the key can actually use."""
    if not key:
        return {"ok": False, "error": "Paste your OpenAI key first."}
    try:
        r = http.get(f"{OPENAI}/models",
                     headers={"Authorization": f"Bearer {key}"}, timeout=20)
    except Exception as e:
        return {"ok": False, "error": f"Could not reach OpenAI ({type(e).__name__})."}
    if r.status_code == 401:
        return {"ok": False, "error": "OpenAI rejected that key."}
    if r.status_code != 200:
        return {"ok": False,
                "error": f"OpenAI said {r.status_code}: {r.text[:160]}"}
    try:
        ids = {m["id"] for m in r.json().get("data", [])}
    except Exception:
        ids = set()
    usable = [m["id"] for m in MODELS if m["id"] in ids] or [DEFAULT_MODEL]
    return {"ok": True, "models": usable}


def complete(key, model, system, user, max_tokens=700, json_mode=False,
             timeout=90):
    """-> {ok, text, data, tokens_in, tokens_out, error}"""
    if not key:
        return {"ok": False, "error": "No OpenAI key on this account."}
    body = {"model": model or DEFAULT_MODEL,
            "messages": [{"role": "system", "content": system},
                         {"role": "user", "content": user}],
            "max_tokens": max_tokens, "temperature": 0.7}
    if json_mode:
        body["response_format"] = {"type": "json_object"}
    try:
        r = http.post(f"{OPENAI}/chat/completions",
                      headers={"Authorization": f"Bearer {key}"},
                      json=body, timeout=timeout)
    except Exception as e:
        return {"ok": False, "error": f"OpenAI did not answer ({type(e).__name__})."}
    if r.status_code == 429:
        return {"ok": False, "retry": True,
                "error": "OpenAI is rate-limiting this key. Slowing down."}
    if r.status_code != 200:
        msg = r.text[:200]
        try:
            msg = r.json().get("error", {}).get("message", msg)
        except Exception:
            pass
        return {"ok": False, "error": f"OpenAI said {r.status_code}: {msg}"}
    d = r.json()
    text = (d.get("choices") or [{}])[0].get("message", {}).get("content", "")
    usage = d.get("usage") or {}
    out = {"ok": True, "text": (text or "").strip(),
           "tokens_in": usage.get("prompt_tokens", 0),
           "tokens_out": usage.get("completion_tokens", 0)}
    if json_mode:
        out["data"] = parse_json(text)
    return out


def parse_json(text):
    """Forgiving: strips code fences, finds the outermost object."""
    t = (text or "").strip()
    t = re.sub(r"^```(?:json)?\s*|\s*```$", "", t, flags=re.S)
    try:
        return json.loads(t)
    except ValueError:
        pass
    i, j = t.find("{"), t.rfind("}")
    if i >= 0 and j > i:
        try:
            return json.loads(t[i:j + 1])
        except ValueError:
            return None
    return None


# -------------------------------------------------------------- cost model
def price(model, tokens_in, tokens_out):
    m = MODEL_BY_ID.get(model) or MODEL_BY_ID[DEFAULT_MODEL]
    return round((tokens_in / 1_000_000) * m["in"]
                 + (tokens_out / 1_000_000) * m["out"], 6)


def tokens_of(text):
    """Rough but honest: about four characters a token for English prose.
    Used only for the before-you-run estimate; real usage is billed from
    OpenAI's own numbers."""
    return max(1, math.ceil(len(str(text or "")) / 4))


def estimate(model, rows, prompt_chars, extra_in_chars=0, out_tokens=220,
             variants=1):
    """-> {per_row, total, tokens_in, tokens_out, model, rows}"""
    t_in = tokens_of("x" * (prompt_chars + extra_in_chars))
    t_out = out_tokens * max(1, variants)
    per = price(model, t_in, t_out)
    return {"model": model, "rows": rows, "variants": variants,
            "tokens_in": t_in, "tokens_out": t_out,
            "per_row": round(per, 6), "total": round(per * max(0, rows), 4)}


def money(v):
    v = float(v or 0)
    if v == 0:
        return "$0.00"
    if v < 0.01:
        return f"${v:.4f}"
    return f"${v:,.2f}"


# ------------------------------------------------------------ spreadsheets
MAX_ROWS = 20000


def read_sheet(data, filename=""):
    """CSV or XLSX bytes -> (headers, rows as list-of-dict). Raises ValueError
    with a sentence a non-technical person can act on."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(data)
    return _read_csv(data)


def _read_csv(data):
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("That file isn't readable as text. Save it as CSV or "
                         "XLSX and try again.")
    sample = text[:8000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = list(reader)
    if not rows:
        raise ValueError("That file is empty.")
    return _shape(rows)


def _read_xlsx(data):
    try:
        import openpyxl
    except ImportError:
        raise ValueError("Save it as CSV — this server can't read XLSX.")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True,
                                    data_only=True)
    except Exception:
        raise ValueError("That spreadsheet wouldn't open. Try saving it as CSV.")
    ws = wb[wb.sheetnames[0]]
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(["" if c is None else str(c) for c in r])
        if len(rows) > MAX_ROWS + 1:
            break
    wb.close()
    if not rows:
        raise ValueError("That spreadsheet's first sheet is empty.")
    return _shape(rows)


def _shape(rows):
    header = [str(h or "").strip() for h in rows[0]]
    seen, headers = {}, []
    for i, h in enumerate(header):
        h = h or f"Column {i + 1}"
        if h in seen:
            seen[h] += 1
            h = f"{h} ({seen[h]})"
        else:
            seen[h] = 1
        headers.append(h)
    out = []
    for raw in rows[1:]:
        if not any(str(c).strip() for c in raw):
            continue
        d = {}
        for i, h in enumerate(headers):
            d[h] = str(raw[i]).strip() if i < len(raw) and raw[i] is not None else ""
        out.append(d)
        if len(out) >= MAX_ROWS:
            break
    if not out:
        raise ValueError("That file has a header row but no data under it.")
    return headers, out


def write_csv(headers, rows):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    for r in rows:
        w.writerow([_safe(r.get(h, "")) for h in headers])
    return buf.getvalue().encode("utf-8-sig")


def _safe(v):
    """Excel runs a cell that starts with = + - @ as a formula. A company name
    scraped off the web must never execute on someone's laptop."""
    s = "" if v is None else str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@", "\t", "\r") else s


# ------------------------------------------------------------------- misc
TOKEN_RE = re.compile(r"\{\{?\s*([A-Za-z0-9 _./#'()&-]+?)\s*\}?\}")


def fill(template, row):
    """Replace {Column} / {{Column}} with that row's value. Unknown tokens are
    left visible on purpose -- a silent blank is how a mail merge ships
    'Hi ,' to four hundred people."""
    missing = []

    def sub(m):
        key = m.group(1).strip()
        for k in (key, key.lower(), key.title(), key.upper()):
            if k in row:
                return str(row[k] or "")
        for k, v in row.items():
            if k.lower().replace(" ", "") == key.lower().replace(" ", ""):
                return str(v or "")
        missing.append(key)
        return "{" + key + "}"

    return TOKEN_RE.sub(sub, template or ""), missing


def tokens_in_template(template):
    return [m.group(1).strip() for m in TOKEN_RE.finditer(template or "")]
